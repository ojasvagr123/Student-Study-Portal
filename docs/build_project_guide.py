from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.style import WD_STYLE_TYPE
from pathlib import Path

OUT = Path(__file__).resolve().parent / "Student_Study_Portal_Complete_Learning_and_Interview_Guide.docx"
NAVY = "17365D"; BLUE = "2E74B5"; DARK = "1F4D78"; PALE = "E8EEF5"; LIGHT = "F4F6F9"; GOLD = "9C6500"; RED = "9B1C1C"; GREEN = "2E6B3B"; GRAY = "5B6573"; WHITE = "FFFFFF"

doc = Document()
sec = doc.sections[0]
sec.top_margin = sec.bottom_margin = sec.left_margin = sec.right_margin = Inches(1)
sec.header_distance = sec.footer_distance = Inches(.492)

styles = doc.styles
normal = styles['Normal']; normal.font.name = 'Calibri'; normal.font.size = Pt(11); normal.font.color.rgb = RGBColor.from_string('1F2937')
normal.paragraph_format.space_after = Pt(6); normal.paragraph_format.line_spacing = 1.25
for name, size, color, before, after in [('Heading 1',16,BLUE,18,10),('Heading 2',13,BLUE,14,7),('Heading 3',12,DARK,10,5)]:
    s=styles[name]; s.font.name='Calibri'; s.font.size=Pt(size); s.font.bold=True; s.font.color.rgb=RGBColor.from_string(color); s.paragraph_format.space_before=Pt(before); s.paragraph_format.space_after=Pt(after); s.paragraph_format.keep_with_next=True
for name in ['List Bullet','List Number']:
    s=styles[name]; s.font.name='Calibri'; s.font.size=Pt(11); s.paragraph_format.left_indent=Inches(.375); s.paragraph_format.first_line_indent=Inches(-.188); s.paragraph_format.space_after=Pt(4); s.paragraph_format.line_spacing=1.25
for name, base in [('Code Block','Normal'),('Callout','Normal'),('Table Text','Normal')]:
    if name not in styles:
        styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
styles['Code Block'].font.name='Consolas'; styles['Code Block'].font.size=Pt(8.5); styles['Code Block'].font.color.rgb=RGBColor.from_string('243447'); styles['Code Block'].paragraph_format.space_after=Pt(8); styles['Code Block'].paragraph_format.line_spacing=1.0
styles['Callout'].font.name='Calibri'; styles['Callout'].font.size=Pt(10.5); styles['Callout'].paragraph_format.left_indent=Inches(.15); styles['Callout'].paragraph_format.right_indent=Inches(.15); styles['Callout'].paragraph_format.space_before=Pt(5); styles['Callout'].paragraph_format.space_after=Pt(8)
styles['Table Text'].font.name='Calibri'; styles['Table Text'].font.size=Pt(9); styles['Table Text'].paragraph_format.space_after=Pt(2); styles['Table Text'].paragraph_format.line_spacing=1.08

def shade(cell, fill):
    tcPr=cell._tc.get_or_add_tcPr(); shd=tcPr.find(qn('w:shd'))
    if shd is None: shd=OxmlElement('w:shd'); tcPr.append(shd)
    shd.set(qn('w:fill'),fill)
def cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc=cell._tc; tcPr=tc.get_or_add_tcPr(); tcMar=tcPr.first_child_found_in('w:tcMar')
    if tcMar is None: tcMar=OxmlElement('w:tcMar'); tcPr.append(tcMar)
    for m,v in [('top',top),('start',start),('bottom',bottom),('end',end)]:
        node=tcMar.find(qn('w:'+m))
        if node is None: node=OxmlElement('w:'+m); tcMar.append(node)
        node.set(qn('w:w'),str(v)); node.set(qn('w:type'),'dxa')
def set_cell_width(cell, width):
    tcPr=cell._tc.get_or_add_tcPr(); tcW=tcPr.find(qn('w:tcW'))
    if tcW is None: tcW=OxmlElement('w:tcW'); tcPr.append(tcW)
    tcW.set(qn('w:w'),str(width)); tcW.set(qn('w:type'),'dxa')
def table(rows, widths):
    t=doc.add_table(rows=1, cols=len(rows[0])); t.alignment=WD_TABLE_ALIGNMENT.LEFT; t.autofit=False
    tblPr=t._tbl.tblPr; tblW=tblPr.find(qn('w:tblW')); tblW.set(qn('w:w'),'9360'); tblW.set(qn('w:type'),'dxa')
    tblInd=OxmlElement('w:tblInd'); tblInd.set(qn('w:w'),'120'); tblInd.set(qn('w:type'),'dxa'); tblPr.append(tblInd)
    grid=t._tbl.tblGrid
    for child in list(grid): grid.remove(child)
    for w in widths:
        gc=OxmlElement('w:gridCol'); gc.set(qn('w:w'),str(w)); grid.append(gc)
    for i,val in enumerate(rows[0]):
        c=t.rows[0].cells[i]; set_cell_width(c,widths[i]); shade(c,PALE); cell_margins(c); c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p=c.paragraphs[0]; p.style='Table Text'; r=p.add_run(str(val)); r.bold=True; r.font.color.rgb=RGBColor.from_string(NAVY)
    for row in rows[1:]:
        cells=t.add_row().cells
        for i,val in enumerate(row):
            c=cells[i]; set_cell_width(c,widths[i]); cell_margins(c); c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p=c.paragraphs[0]; p.style='Table Text'; p.add_run(str(val))
    doc.add_paragraph().paragraph_format.space_after=Pt(1)
    return t
def p(text='', bold_lead=None):
    q=doc.add_paragraph()
    if bold_lead and text.startswith(bold_lead):
        q.add_run(bold_lead).bold=True; q.add_run(text[len(bold_lead):])
    else: q.add_run(text)
    return q
def bullet(text): doc.add_paragraph(text, style='List Bullet')
def num(text): doc.add_paragraph(text, style='List Number')
def code(text):
    q=doc.add_paragraph(style='Code Block'); q.paragraph_format.left_indent=Inches(.18); q.paragraph_format.right_indent=Inches(.12); shade_paragraph(q,'F3F6F8'); q.add_run(text)
def shade_paragraph(q, fill):
    pPr=q._p.get_or_add_pPr(); shd=OxmlElement('w:shd'); shd.set(qn('w:fill'),fill); pPr.append(shd)
def callout(label,text,color=BLUE,fill=LIGHT):
    q=doc.add_paragraph(style='Callout'); shade_paragraph(q,fill); r=q.add_run(label+' '); r.bold=True; r.font.color.rgb=RGBColor.from_string(color); q.add_run(text)
def heading(text,level=1): doc.add_heading(text,level=level)
def page(): doc.add_page_break()

# Running furniture
h=sec.header.paragraphs[0]; h.text='STUDENT STUDY PORTAL  |  PROJECT LEARNING GUIDE'; h.style='Table Text'; h.runs[0].font.color.rgb=RGBColor.from_string(GRAY)
f=sec.footer.paragraphs[0]; f.alignment=WD_ALIGN_PARAGRAPH.RIGHT; f.add_run('Interview preparation edition  |  ')
fld=OxmlElement('w:fldSimple'); fld.set(qn('w:instr'),'PAGE'); f._p.append(fld)

# Cover
doc.add_paragraph().paragraph_format.space_after=Pt(80)
q=doc.add_paragraph(); q.alignment=WD_ALIGN_PARAGRAPH.CENTER; r=q.add_run('COMPLETE PROJECT GUIDE'); r.bold=True; r.font.name='Calibri'; r.font.size=Pt(12); r.font.color.rgb=RGBColor.from_string(GOLD)
q=doc.add_paragraph(); q.alignment=WD_ALIGN_PARAGRAPH.CENTER; q.paragraph_format.space_after=Pt(10); r=q.add_run('Student Study Portal'); r.bold=True; r.font.size=Pt(30); r.font.color.rgb=RGBColor.from_string(NAVY)
q=doc.add_paragraph(); q.alignment=WD_ALIGN_PARAGRAPH.CENTER; q.paragraph_format.space_after=Pt(30); r=q.add_run('Learn the codebase from first principles to deployment - and defend every decision in an interview'); r.italic=True; r.font.size=Pt(14); r.font.color.rgb=RGBColor.from_string(DARK)
table([['Stack','Architecture','Deployment story'],['React 19 + Tailwind','Spring Boot modular monolith','Vercel + Render + Supabase'],['Java 17 + Spring Boot 3.1','REST + JWT + JPA','Repository currently needs production config']], [3120,3120,3120])
callout('Purpose.', 'Use this as a reading path, code walkthrough, architecture reference, system-design case study, and interview workbook. It distinguishes verified code from deployment claims that are not reproducible from the current repository.', NAVY, 'EEF3F8')
p('Prepared from the repository snapshot on 21 August 2026. This is an implementation audit, not a generic Spring Boot tutorial.')
page()

heading('How to use this guide',1)
callout('Recommended path.', 'Read Chapters 1-3 for the two-minute project story; Chapters 4-10 while opening the referenced files; Chapters 11-14 for production and system-design depth; then answer Chapter 15 aloud without looking.')
table([['Pass','Focus','Outcome'],['1. Orientation (2-3 hours)','Product, features, stack, architecture','Explain the project in 90 seconds'],['2. Code tracing (1-2 days)','Auth, academic resource, career, contribution flows','Trace UI -> API -> database'],['3. Engineering depth (1 day)','Security, JPA, trade-offs, patterns, failures','Defend technology choices'],['4. Interview rehearsal (1 day)','Questions, model answers, improvements','Speak honestly and confidently']], [1300,3400,4660])
heading('Repository truth labels',2)
bullet('VERIFIED: directly present in source/configuration and described as implemented.')
bullet('INFERRED: a reasonable architectural interpretation, explicitly labeled.')
bullet('TARGET: recommended production design or improvement, not claimed as current code.')
bullet('DEPLOYMENT CLAIM: Render/Vercel/Supabase were stated by the project owner, but the current repository has no deployment manifests or production environment variables to reproduce it.')
heading('Contents',2)
for x in ['1. Product and feature map','2. Technology choices and alternatives','3. High-level architecture and system design','4. Repository map and startup','5. Backend architecture and Spring lifecycle','6. Authentication and authorization walkthrough','7. Academic resource flow','8. Career, quiz, progress, contribution, comment, todo flows','9. Database and JPA model','10. Frontend architecture and request flow','11. Design patterns','12. Deployment: Vercel, Render, Supabase','13. Testing, quality, security and production gaps','14. Scaling and HLD evolution','15. Interview preparation','16. Learning plan and cheat sheets','17. Future product and AI integration blueprint']:
    bullet(x)
page()

heading('1. Product and feature map',1)
p('Student Study Portal is a full-stack student companion. It combines semester-based academic resources with career preparation and community participation. The application is best described as a REST-based modular monolith: one React single-page application, one Spring Boot backend process, and one PostgreSQL database.')
heading('Implemented feature inventory',2)
table([['Module','User value','Primary code'],['Authentication','Register, login, JWT-protected identity','AuthController, SecurityConfig, JwtAuthenticationFilter, JwtUtil'],['Academics','Browse semesters, branch subjects, units, resources and PYQs','StudyResourceController/Service and five repositories'],['Career preparation','Goal-filtered resources and placement topic progress','CareerController, CareerResource, StriverTopic'],['Quiz','Random questions by career goal and answer scoring','QuizController, QuizQuestionRepository'],['Community','Submit resources, list own contributions, comment in context','ContributionController, CommentController'],['Productivity','Todos, weekly schedule, reminders/calendar UI','TodoController and React components'],['Profile','Authenticated profile lookup and overview UI','UserController, Profile components']], [1800,3500,4060])
heading('What the project is not',2)
bullet('It is not microservices: all backend modules deploy together and share one database.')
bullet('It is not server-side rendered: React renders in the browser and calls JSON APIs.')
bullet('It is not fully production-configured in this snapshot: API URLs, CORS origins, DB credentials, and JWT secret are hard-coded for local development.')
bullet('It is not comprehensively tested: the backend build reports NO-SOURCE for tests; the frontend retains only the Create React App sample test.')
heading('A strong 90-second interview introduction',2)
callout('Model answer.', '“I built Student Study Portal, a full-stack platform that centralizes academic resources and career preparation. The frontend is a React SPA using React Router, Axios, Zustand, Tailwind and feature components. The backend is a Java 17 Spring Boot modular monolith with REST controllers, Spring Security, stateless JWT authentication, BCrypt password hashing, Spring Data JPA and PostgreSQL. Core flows include semester-to-unit content discovery, career resources, quizzes, placement progress, contributions, comments and todos. I deployed the frontend, backend and managed PostgreSQL separately using Vercel, Render and Supabase. Reviewing it later, I identified production gaps such as hard-coded URLs/secrets, broad authorization rules, local file storage and missing automated tests; I can explain how I would fix each.”')

heading('2. Technology choices and alternatives',1)
table([['Choice','Why it fits','Why not the common alternative?'],['Java 17','LTS-era language, strong typing, mature tooling, Spring Boot compatibility','Node/Express is faster to prototype but offers fewer compile-time domain guarantees; Django is productive but changes the language/ecosystem.'],['Spring Boot 3.1','Auto-configuration, embedded server, production conventions, integrated web/security/data ecosystem','Plain Spring requires more configuration; Jakarta EE servers add deployment overhead; microframeworks trade ecosystem depth for smaller footprint.'],['Spring MVC / REST','Natural CRUD/resource APIs, easy React integration, stateless HTTP','GraphQL is valuable for client-driven aggregation but adds schema/resolver/cache complexity not required here.'],['Spring Data JPA / Hibernate','Repository abstraction, derived queries, entity relationships, reduced JDBC boilerplate','Raw JDBC gives SQL control but adds mapping and transaction boilerplate; jOOQ is excellent for SQL-first systems but was unnecessary at this scale.'],['PostgreSQL','Relational integrity, joins, transactions, arrays in the proposed schema, managed hosting on Supabase','MongoDB is flexible but the domain has strong relations; MySQL would work, but PostgreSQL offers richer types and strong hosted options.'],['JWT + BCrypt','Stateless API authentication and industry-standard adaptive password hashing','Server sessions simplify revocation but need shared session state when scaled; OAuth/OIDC is better for social/enterprise identity but more complex.'],['React','Component model, ecosystem, SPA routing and fast UI iteration','Angular is more opinionated/heavier; Vue is simpler but was not the selected ecosystem; server templates reduce frontend separation.'],['Zustand','Small global auth store with minimal ceremony','Redux is powerful for complex event/state flows but excessive for the current state footprint; Context can cause broad rerenders and scattered logic.'],['Axios','Interceptors and ergonomic JSON/error handling','fetch is built-in but needs repeated wrappers for base URL, headers and errors.'],['Tailwind CSS','Rapid utility-based responsive styling and consistent tokens','CSS modules avoid utility-heavy markup but require more stylesheet navigation; component libraries can constrain custom design.'],['Vercel','Excellent static SPA build/CDN workflow','Serving React from Spring simplifies deployment but couples release cycles and loses edge CDN convenience.'],['Render','Simple managed Java web service and Git-based deployment','A VM gives control but adds operations; Kubernetes is disproportionate for this application.'],['Supabase PostgreSQL','Managed Postgres, connection details, backups/dashboard','Self-hosted Postgres adds patching, backup and availability work; Supabase Auth was not needed because Spring Security owns authentication.']], [1700,3640,4020])
callout('Interview principle.', 'Do not say a technology is universally “best.” State the requirement, why the choice fit this project, the trade-off accepted, and the condition under which you would choose another option.', GOLD, 'FFF8E8')

heading('3. High-level architecture and system design',1)
heading('Context diagram',2)
code('Student browser\n    | HTTPS\n    v\n[Vercel: React static SPA]\n    | JSON/HTTPS + Authorization: Bearer <JWT>\n    v\n[Render: Spring Boot REST API]\n    | JDBC/TLS\n    v\n[Supabase: managed PostgreSQL]\n\nExternal links: YouTube/resource URLs, Google Calendar template URL\nCurrent upload implementation: Render-local /uploads directory (ephemeral risk)')
heading('Container responsibilities',2)
bullet('React SPA: routes, presentation, form state, local token storage, Axios calls, client-side protected routes.')
bullet('Spring Boot API: validation, authentication, authorization, orchestration, persistence mapping, file upload endpoints, JSON responses.')
bullet('PostgreSQL: durable relational state for users, academics, career data, contributions, comments, quiz questions, progress and todos.')
heading('Backend layered view',2)
code('HTTP request\n  -> SecurityFilterChain\n  -> JwtAuthenticationFilter\n  -> @RestController\n  -> @Service (StudyResourceService where present)\n  -> JpaRepository proxy\n  -> Hibernate/JPA\n  -> PostgreSQL\n  -> entity/DTO serialized to JSON\n  -> React state -> component render')
heading('Key quality attributes',2)
table([['Attribute','Current approach','Main limitation'],['Security','BCrypt, signed JWT, stateless filter, method annotations','Fallback permitAll, hard-coded secrets, inconsistent authority naming, localStorage XSS exposure'],['Maintainability','Packages by layer, DTOs, repository interfaces','Many controllers call repositories directly; field injection; duplicate frontend entry/auth/API code'],['Scalability','Stateless API can be horizontally replicated','Local uploads, no caching/pagination, DB connection limits, Render cold starts'],['Reliability','Managed DB target, clear CRUD flow','ddl-auto=update, no migrations, weak error strategy, no health/observability configuration'],['Performance','JPA query derivation and a targeted JPQL query','Potential N+1/recursive serialization, random-order query cost, unbounded list endpoints'],['Deployability','Independent frontend/backend/database targets','No deployment manifests and localhost values in tracked configuration']], [1700,3800,3860])
heading('Principal data flows',2)
num('Login: credentials -> AuthenticationManager -> DaoAuthenticationProvider -> UserDetailsService -> BCrypt comparison -> JwtUtil -> token returned -> browser localStorage.')
num('Academic browse: route semester -> GET subjects with branch -> JPQL over Subject/Branch -> choose subject -> load Units -> choose unit -> load Resources and PYQs.')
num('Authenticated contribution: Axios interceptor adds bearer token -> JWT filter builds SecurityContext -> controller reads principal -> saves Contribution linked to User.')
num('Quiz: client requests goal -> repository selects matching questions in random order -> client submits answer map -> controller calculates score -> response updates UI.')

heading('4. Repository map and startup',1)
code('Student-Study-Portal/\n  backend/\n    build.gradle                 Spring Boot/Java dependencies\n    src/main/java/com/studentcompanion/\n      config/                    security, CORS, MVC static uploads, seeding\n      controller/                HTTP endpoints\n      dto/                       API request/response shapes\n      filter/                    JWT request filter\n      model/                     JPA entities and enums\n      repository/                Spring Data interfaces\n      service/                   domain/application services\n      util/                      JWT creation and parsing\n    src/main/resources/application.properties\n  frontend/\n    package.json                 React ecosystem\n    src/App.js                   route composition\n    src/components/              screens and reusable pieces\n    src/pages/                   career/profile/contribution pages\n    src/stores/authStore.js      Zustand auth state\n    src/utils/axiosConfig.js     shared HTTP client\n  database/schema.sql            ambitious target schema (not fully mapped)\n  backend/schema.sql             seed/data schema artifact')
heading('Local startup sequence',2)
num('Create PostgreSQL database studentcompanion and use a non-production local account.')
num('Replace application.properties secrets with environment variables or a local ignored profile.')
num('From backend, run gradlew.bat bootRun. Spring scans from com.studentcompanion, creates repository proxies, configures Hibernate/Security, and starts embedded Tomcat on port 8080.')
num('From frontend, run npm start. Create React App serves port 3000 and the browser calls the backend.')
num('Register, login, preserve the returned token, then inspect Network requests and backend SQL logs while navigating.')
code('# Safer configuration pattern\nspring.datasource.url=${SPRING_DATASOURCE_URL}\nspring.datasource.username=${SPRING_DATASOURCE_USERNAME}\nspring.datasource.password=${SPRING_DATASOURCE_PASSWORD}\njwt.secret=${JWT_SECRET}\napp.cors.allowed-origins=${APP_CORS_ALLOWED_ORIGINS:http://localhost:3000}')
callout('Verified build result.', 'On this repository snapshot, Gradle compilation succeeds. There are no backend test sources, so Gradle reports “test NO-SOURCE.” This proves compilation, not behavioral correctness.', GREEN, 'EDF7EF')

heading('5. Backend architecture and Spring lifecycle',1)
heading('Bootstrap and dependency graph',2)
p('StudentCompanionApplication carries @SpringBootApplication, which combines configuration, component scanning and auto-configuration. @EnableJpaAuditing activates auditing infrastructure, although the entity audit usage should be checked. @EnableJpaRepositories is explicit but redundant because repository packages are already under the application root package.')
code('@SpringBootApplication\n@EnableJpaAuditing\n@EnableJpaRepositories(basePackages = "com.studentcompanion.repository")\npublic class StudentCompanionApplication {\n  public static void main(String[] args) {\n    SpringApplication.run(StudentCompanionApplication.class, args);\n  }\n}')
heading('Controllers',2)
p('Controllers translate HTTP into Java calls. @RequestMapping defines a resource prefix; method annotations define verbs and paths; @RequestBody binds JSON; @PathVariable and @RequestParam bind URL input; ResponseEntity controls status and body. The repository contains controllers for auth, users, academics, career, quizzes, contributions, comments, todos, admin/test/internal endpoints and a legacy HTML-style HomeController.')
heading('Service layer',2)
p('StudyResourceService is the clearest application-service example: it coordinates semester, subject, unit, resource and PYQ repositories. Other controllers frequently use repositories directly. That is acceptable for very small CRUD operations but creates inconsistent business boundaries. A production refactor would introduce AuthService, ContributionService, CommentService, CareerService and QuizService, each transactional and tested.')
heading('Repository layer',2)
p('JpaRepository<Entity, Long> supplies CRUD, paging and sorting. Spring parses method names such as findByUnitId and generates implementations at runtime. SubjectRepository uses JPQL because it needs a branch-code/semester predicate across a relationship. QuizQuestionRepository calls PostgreSQL RANDOM through JPQL function integration; simple, but expensive for a large table.')
code('public interface UnitRepository extends JpaRepository<Unit, Long> {\n  List<Unit> findBySubjectId(Long subjectId);\n}\n\n@Query("SELECT s FROM Subject s WHERE s.branch.code = :code AND s.semester = :semester")\nList<Subject> getSubjectsByBranchAndSemester(String code, int semester);')
heading('DTOs versus entities',2)
p('The project defines request/response DTOs for authentication, registration, comments, contributions, quizzes and study resources. DTOs protect the HTTP contract from persistence details and prevent accidental password/relationship serialization. However, several endpoints still return entities directly. Standardize on DTOs plus mapping functions; never serialize User.password.')
heading('Dependency injection improvement',2)
code('// Current pattern: mutable field injection\n@Autowired private UserRepository userRepository;\n\n// Preferred: immutable constructor injection\nprivate final UserRepository users;\npublic AuthService(UserRepository users) { this.users = users; }')
bullet('Constructor injection makes dependencies explicit and enables unit tests without reflection.')
bullet('Use @Transactional on service methods that update multiple rows as one business operation.')
bullet('Use a @RestControllerAdvice for consistent ErrorResponse mapping instead of scattered try/catch blocks.')

heading('6. Authentication and authorization walkthrough',1)
heading('Registration',2)
num('React SignupPage validates form fields and POSTs registration JSON.')
num('AuthController validates input/email and checks whether the email exists.')
num('PasswordEncoder hashes the password using BCrypt; the plaintext should never be stored or logged.')
num('A User is saved with a Role; the response should contain safe user data and optionally a JWT.')
heading('Login',2)
num('AuthController passes email/password to AuthenticationManager.')
num('DaoAuthenticationProvider asks CustomUserDetailsService to load the database user by email.')
num('BCryptPasswordEncoder compares submitted password with stored hash.')
num('JwtUtil signs an HS256 token whose subject is the email and whose expiry is configured in seconds.')
num('The frontend stores user/token and sends Authorization: Bearer token on later requests.')
heading('Authenticated request',2)
code('Authorization: Bearer eyJ...\n  -> JwtAuthenticationFilter (OncePerRequestFilter)\n  -> extract subject and expiry; verify HMAC signature\n  -> load UserDetails again\n  -> UsernamePasswordAuthenticationToken\n  -> SecurityContextHolder\n  -> controller can use Authentication.getName()')
heading('Security configuration explained',2)
bullet('CSRF is disabled because the API is intended to be stateless and uses an Authorization header, not an automatically attached session cookie.')
bullet('SessionCreationPolicy.STATELESS prevents server-side login session storage.')
bullet('The JWT filter runs before UsernamePasswordAuthenticationFilter so the SecurityContext is ready for authorization.')
bullet('@EnableMethodSecurity enables @PreAuthorize checks on admin methods.')
heading('Critical security findings',2)
table([['Finding','Why it matters','Fix'],['Plain DB password and JWT secret in tracked properties','Credential compromise and token forgery if leaked','Rotate both; environment variables/secret manager; purge historical secrets if published'],['.anyRequest().permitAll()','Endpoints not explicitly matched are public','Default to authenticated; explicitly permit only public routes'],['CORS only localhost','Deployed Vercel SPA will be blocked','Environment-driven exact production origin; no wildcard with credentials'],['@CrossOrigin("*") plus central CORS','Conflicting policy and overly broad intent','Remove controller annotations; manage one policy centrally'],['Role/authority inconsistency','hasRole adds ROLE_ prefix; stored authority appears ADMIN','Adopt ROLE_ADMIN everywhere or use hasAuthority consistently'],['Token in localStorage','Any XSS can steal bearer token','Harden CSP/input handling; consider short-lived access token plus secure HttpOnly refresh cookie'],['No refresh/revocation','Stolen token valid until expiry','Short expiry, refresh rotation, logout deny-list only if threat model requires'],['Internal promotion endpoint','Privilege-escalation risk if reachable','Remove from production or strongly restrict operationally']], [2000,3400,3960])
callout('Interview answer.', 'JWT is not encryption; its payload is readable. Security comes from signature verification, short expiry, safe key management and TLS. BCrypt protects stored passwords through salted adaptive hashing.', RED, 'FCEEEE')

heading('7. Academic resource flow',1)
heading('Domain hierarchy',2)
code('Branch\n  -> Subject (semester + branch)\n       -> Unit\n            -> Resource (notes/video/link/file)\n            -> PYQ (previous-year question)\n                 -> PyqComment (separate model)')
heading('API journey',2)
table([['Step','HTTP API','Backend path'],['Load semester catalog','GET /api/semesters','StudyResourceController -> StudyResourceService -> SemesterRepository.findAll'],['Filter subjects','GET /api/semesters/{semester}/subjects?branch=CSE','SubjectRepository JPQL filters branch.code and semester'],['Load units','GET /api/subjects/{id}/units','UnitRepository.findBySubjectId'],['Load unit details','GET /api/units/{id}','Unit lookup in controller/repository'],['Load learning material','GET /api/units/{id}/resources','ResourceRepository.findByUnitId'],['Load past questions','GET /api/units/{id}/pyqs','PyqRepository.findByUnitId'],['Admin creation/upload','POST /api/admin/...','@PreAuthorize ADMIN methods']], [1800,3600,3960])
heading('Frontend trace',2)
p('SubjectsPage reads semesterId from the URL and currently supplies a fixed/default branch value when requesting subjects. UnitPage loads all units for a subject, tracks the active unit index, then loads resources and PYQs for the selected unit. This is a master-detail UI pattern.')
heading('JPA concerns to understand',2)
bullet('Bidirectional associations can recurse during JSON serialization. Prefer DTOs or @JsonIgnore/managed references deliberately.')
bullet('LAZY relationships save queries until accessed; accessing after the transaction closes can raise LazyInitializationException.')
bullet('EAGER loading can cause large joins or N+1 behavior. Fetch only what each API requires using DTO projection/entity graphs.')
bullet('Foreign-key indexes should support unit_id, subject_id, branch_id+semester and user-owned lookups.')
heading('Admin upload concern',2)
p('WebConfig exposes files from user.dir/uploads. This works locally. Render instances use an ephemeral filesystem and may restart or scale to multiple instances, causing files to disappear or diverge. TARGET: upload to object storage (Supabase Storage/S3), persist a durable URL and validate MIME type, size, filename, ownership and authorization.')

heading('8. Remaining feature flows',1)
heading('Career resources',2)
p('CareerController maps a CareerGoal enum and optional ResourceType to repository filters. CareerGoalPage reads /goal/:goal from React Router, converts the goal for the backend enum, fetches resources and composes progress/quiz UI. Enums constrain allowed categories but require stable API naming or explicit conversion.')
heading('Placement progress / Striver topics',2)
p('The placement endpoints retrieve topics belonging to the authenticated user and mark attempts. The natural uniqueness rule is (user_id, topic). StriverTopicRepository exposes findByUser and findByUserAndTopic, allowing an idempotent upsert-like flow rather than duplicate progress rows.')
heading('Quiz',2)
p('QuizQuestionRepository retrieves goal-specific questions ordered randomly. QuizController exposes GET /api/quiz/{goal} and POST /api/quiz/submit. QuizSubmissionDTO carries submitted answers; scoring belongs on the server so a client cannot declare its own score. Do not return correct answers before submission in a real assessment.')
heading('Contribution and comments',2)
p('ContributionController creates a Contribution attached to the authenticated user and lists “my” contributions. CommentController stores a generic contextType + contextId, which is a polymorphic association at application level: a comment can target multiple resource kinds without one foreign key per type.')
callout('Trade-off.', 'The generic comment target is flexible, but PostgreSQL cannot enforce that contextId points to a real row of the declared type. Validate allowed types and target existence in a service; a dedicated table per target gives stronger integrity but more schema/API complexity.', GOLD, 'FFF8E8')
heading('Todos',2)
p('TodoController provides list-by-user, create, update and delete operations. The security configuration currently permits /api/todos/**, and an API accepting arbitrary userId is vulnerable to insecure direct object reference. Derive the owner from Authentication, not from a caller-controlled ID, and verify ownership on update/delete.')
heading('Frontend/backend contract mismatches found',2)
table([['Mismatch','Evidence','Effect'],['Auth store hydration','Zustand initializes token=null even when localStorage has a token; ProtectedRoute reads localStorage while Axios reads Zustand','Page refresh may pass route protection but omit Authorization header'],['setUser signature','auth.js calls setUser(user) although store expects (user, token)','Token becomes undefined'],['User comments route','Frontend requests /api/comments/user/{user.id}; controller exposes GET /api/comments/user','404 or contract drift'],['Profile edit route','Frontend PUT /api/users/{id}; shown UserController exposes all/profile only','Update fails unless route exists elsewhere'],['Calendar assignments','Frontend calls /api/calendar/assignments; no matching controller found','Feature is UI-only/incomplete'],['API base URLs','Several components hard-code localhost while others use shared Axios','Production deployment breaks inconsistently'],['Entry points/providers','Both index.js and main.jsx exist; ThemeProvider is nested in main/App','Dead/duplicate boot path and redundant context']], [1800,3200,4360])

heading('9. Database and JPA model',1)
heading('Implemented entity families',2)
table([['Family','Entities/enums','Relationship intent'],['Identity','User, Role, Branch, CareerGoal','User owns profile/preferences and feature records'],['Academics','Semester, Subject, Unit, Resource, ResourceType, Pyq, PyqComment','Hierarchical curriculum and material'],['Career','CareerResource, QuizQuestion, StriverTopic','Goal/type resources, assessment and per-user progress'],['Community','Contribution, Comment','User-generated resources and generic contextual discussion'],['Productivity','Todo','Per-user task tracking']], [1600,3300,4460])
heading('Schema reality versus design artifact',2)
p('database/schema.sql describes a broader future platform: profiles, companies, coding sheets, progress, achievements, discussions, votes and notifications. Many of those tables have no corresponding Java entity/controller. Treat it as a target/brainstorm schema, not the authoritative runtime schema. Hibernate ddl-auto=update currently makes Java entities the practical schema source.')
heading('Why ddl-auto=update is risky',2)
bullet('Schema changes are implicit and environment-dependent; destructive or complex migrations are not safely represented.')
bullet('Teams cannot review a deterministic versioned migration history.')
bullet('TARGET: Flyway or Liquibase migrations, ddl-auto=validate in production, and a clean baseline migration matching entities.')
heading('Transaction and integrity rules to add',2)
bullet('Unique users.email; normalize email case before lookup and store.')
bullet('Unique progress (user_id, topic) and appropriate foreign keys with chosen cascade/delete behavior.')
bullet('Check constraints/enums or lookup tables for status/type values.')
bullet('Optimistic locking (@Version) for concurrently edited todos/progress where lost updates matter.')
bullet('Created/updated timestamps with auditing annotations and database-compatible timezone strategy (UTC).')
heading('Index strategy',2)
table([['Query','Suggested index'],['Subjects for branch and semester','subjects(branch_id, semester) or branch code join support'],['Units/resources/PYQs by parent','units(subject_id), resources(unit_id), pyq(unit_id)'],['User-owned records','todo(user_id), contribution(user_id), comment(user_id)'],['Context comments sorted by time','comment(context_type, context_id, timestamp DESC)'],['Email login','users(lower(email)) UNIQUE'],['Career goal/type filters','career_resource(goal, type), quiz_question(goal)']], [3300,6060])

heading('10. Frontend architecture and request flow',1)
heading('Composition',2)
p('App.js defines BrowserRouter routes and composes ThemeProvider and ToastContainer. ProtectedRoute is a client-side UX gate. Pages/components own local UI state with useState/useEffect; Zustand holds global auth state; Axios handles HTTP. The app uses Tailwind utilities, icons, FullCalendar, Framer Motion and toast notifications.')
heading('Routing map',2)
table([['Route','Screen','Protection'],['/','LandingPage','Public'],['/signup, /login','Account entry','Public'],['/dashboard','DashboardPage','ProtectedRoute'],['/goal/:goal','CareerGoalPage','ProtectedRoute'],['/subjects/:semesterId','SubjectsPage','Client-public currently'],['/subjects/:subjectId/units','UnitPage','Client-public currently'],['/quizzes, /profile, /contribute','Feature pages','Not consistently wrapped']], [2600,4000,2760])
heading('Correct production API client',2)
code('const api = axios.create({\n  baseURL: process.env.REACT_APP_API_BASE_URL,\n});\napi.interceptors.request.use(config => {\n  const token = useAuthStore.getState().token;\n  if (token) config.headers.Authorization = `Bearer ${token}`;\n  return config;\n});')
heading('State and lifecycle improvements',2)
bullet('Hydrate Zustand synchronously from localStorage or use persist middleware so route and HTTP layers agree after refresh.')
bullet('Use one entry point (index.js for Create React App), one ThemeProvider and one shared API client.')
bullet('Separate server state from UI state. TanStack Query would add cache, retries, loading/error normalization and invalidation; it is a future option, not currently installed.')
bullet('Cancel or ignore stale effects to avoid setting state after navigation; expose loading, empty and error states consistently.')
bullet('Never treat ProtectedRoute as authorization. The backend must enforce every protected action.')
heading('Component design discussion',2)
p('The current organization mixes screen-level and reusable components under components/. A scalable structure groups by feature (auth, academics, career, community, profile) and keeps shared/ui separate. This reduces cross-feature coupling and makes tests easier to locate.')
code('src/\n  app/ routes, providers\n  features/auth/\n  features/academics/\n  features/career/\n  features/community/\n  shared/api/ shared/ui/ shared/hooks/')

heading('11. Design patterns in the project',1)
table([['Pattern','Where','Why it matters'],['Layered architecture','Controller -> service -> repository','Separates transport, business orchestration and persistence'],['Repository','JpaRepository interfaces','Collection-like persistence abstraction and generated queries'],['Dependency Injection / IoC','@Component/@Service/@Repository/@Autowired','Spring creates and wires components; code depends on contracts'],['Front Controller','Spring DispatcherServlet (framework-provided)','One HTTP entry dispatches to mapped controller methods'],['Filter chain / Chain of Responsibility','Spring Security filters + JWT filter','Each filter processes/forwards a request in ordered stages'],['Strategy','PasswordEncoder, AuthenticationProvider, UserDetailsService','Algorithms/providers are replaceable behind interfaces'],['DTO','AuthRequest, RegistrationRequest, responses','Stable API boundary and data minimization'],['MVC','React view + Spring MVC controllers (distributed MVC-like split)','Separates representation from server request handling'],['Observer / reactive rendering','React state updates trigger render','UI declaratively reacts to state'],['Singleton scope','Spring beans by default; Zustand store module','Shared application services/state (must remain stateless/thread-safe)'],['Factory/proxy (framework-generated)','Spring repository proxies and bean factory','Runtime implementation of repository interfaces/AOP'],['Polymorphic association','Comment contextType + contextId','One comment model supports multiple target kinds']], [1800,3300,4260])
heading('Patterns not actually present',2)
bullet('Microservices, event sourcing, CQRS, Saga, API gateway and message broker are not in the repository.')
bullet('Do not claim Clean Architecture or Hexagonal Architecture: domain models depend directly on JPA and packages are layer-oriented.')
bullet('Builder/Factory patterns may occur inside libraries, but there is no meaningful custom domain implementation to showcase.')
heading('SOLID evaluation',2)
bullet('SRP: packages help, but large controllers combine validation, authentication, mapping and persistence; extract services.')
bullet('OCP/DIP: Spring interfaces and repositories help, but concrete repositories are injected directly; service ports are limited.')
bullet('ISP: repository interfaces are convenient but expose more CRUD operations than every consumer needs; acceptable here.')

heading('12. Deployment: Vercel + Render + Supabase',1)
callout('Evidence boundary.', 'The owner states these services were used. No vercel.json, render.yaml, Dockerfile, production profile, environment-based API URL or Supabase URL exists in the current repository. The following is the correct reproducible deployment design and likely operational flow, not proof of the historical settings.', RED, 'FCEEEE')
heading('Supabase database',2)
num('Create a Supabase project and obtain the PostgreSQL connection details. Use the direct connection for a long-running Render service when IPv6/network support fits, or the documented pooler endpoint when needed.')
num('Store JDBC URL, username and password only in Render environment secrets. Require SSL according to Supabase settings.')
num('Run versioned migrations (Flyway) before/at deployment. Do not use ddl-auto=update in production.')
num('Size HikariCP below the Supabase connection limit, considering every Render instance.')
code('SPRING_DATASOURCE_URL=jdbc:postgresql://<host>:<port>/<db>?sslmode=require\nSPRING_DATASOURCE_USERNAME=<secret>\nSPRING_DATASOURCE_PASSWORD=<secret>\nSPRING_JPA_HIBERNATE_DDL_AUTO=validate')
heading('Render backend',2)
num('Create a Web Service connected to the backend repository/root.')
num('Build with ./gradlew clean bootJar and start with java -jar build/libs/backend-0.0.1-SNAPSHOT.jar (actual jar name should be verified).')
num('Bind to Render’s PORT: server.port=${PORT:8080}. Configure Java 17.')
num('Set database, JWT_SECRET, allowed Vercel origin and active profile as environment variables.')
num('Expose Spring Boot Actuator health endpoint and configure a health check. This dependency is not currently present.')
code('server.port=${PORT:8080}\napp.cors.allowed-origins=${APP_CORS_ALLOWED_ORIGINS}\njwt.secret=${JWT_SECRET}\nlogging.level.org.springframework.boot.autoconfigure=INFO')
heading('Vercel frontend',2)
num('Import frontend as the project root, install with npm ci and build with npm run build; output is build/.')
num('Set REACT_APP_API_BASE_URL=https://<render-service>.onrender.com before build; Create React App substitutes it at build time.')
num('Configure SPA fallback so deep links return index.html rather than 404.')
num('Add the exact Vercel production and preview origins to backend CORS, or restrict previews deliberately.')
heading('Release request path',2)
code('Developer push\n  -> Vercel builds static React -> CDN\n  -> Render builds bootJar -> Java service\n  -> Spring connects via TLS -> Supabase Postgres\n  -> Browser loads Vercel -> calls Render -> queries Supabase')
heading('Deployment failure checklist',2)
table([['Symptom','Likely cause','Check'],['Browser CORS error','Vercel origin absent or wildcard/credentials conflict','Response Access-Control-Allow-Origin and OPTIONS request'],['Frontend calls localhost','Hard-coded Axios/fetch URL','Build-time env and source-wide URL search'],['Backend cannot start','Bad JDBC URL/SSL/credentials or schema mismatch','Render logs; DB reachability; migration status'],['401 after refresh','Zustand token not hydrated','Store initialization and Authorization request header'],['Uploads disappear','Ephemeral Render disk','Move to object storage'],['Deep-link 404','No SPA rewrite','Vercel rewrite to /index.html'],['Slow first request','Free-tier cold start or DB wake/connect','Health metrics and platform plan']], [1900,3450,4010])

heading('13. Testing, quality and production gaps',1)
heading('Current verification result',2)
callout('Build.', 'gradlew.bat test completed successfully on 21 August 2026 after resolving Gradle dependencies. It compiled main code; compileTestJava and test were NO-SOURCE. Frontend automated verification was not treated as meaningful because App.test.js is the untouched “learn react” sample.', GREEN, 'EDF7EF')
heading('Test pyramid to implement',2)
table([['Layer','Tools','Highest-value cases'],['Unit','JUnit 5, Mockito','Quiz scoring, registration validation, ownership checks, goal/type mapping'],['Web slice','@WebMvcTest, Spring Security test','Status codes, DTO validation, 401/403, JSON contract, admin authorization'],['Repository','@DataJpaTest + Testcontainers PostgreSQL','Derived/JPQL queries, constraints, mappings and migrations'],['Integration','@SpringBootTest + Testcontainers','Register -> login -> authenticated flow; academic browse; contribution/comment'],['Frontend component','React Testing Library + MSW','Forms, route protection, loading/error states, token hydration'],['End-to-end','Playwright/Cypress against staging','Core student journey and deployment wiring']], [1500,2800,5060])
heading('Prioritized engineering backlog',2)
table([['Priority','Change','Reason'],['P0','Rotate/remove tracked DB password and JWT secret','Immediate credential risk'],['P0','Replace permitAll fallback; enforce ownership and admin rules','Broken access control risk'],['P0','Environment-driven API/DB/CORS config','Deployment correctness'],['P1','Unify auth hydration/API client/frontend entry point','Refresh and production request reliability'],['P1','Add Flyway; ddl-auto=validate','Reproducible schema'],['P1','Add exception handler, DTO validation and safe response DTOs','Predictable API/security'],['P1','Object storage for uploads','Durability and horizontal scaling'],['P1','Automated auth/ownership/integration tests','Regression protection'],['P2','Pagination, indexes, query profiling and projections','Scale/performance'],['P2','Observability: Actuator, structured logs, metrics, tracing IDs','Operations'],['P2','Upgrade supported dependency versions after compatibility testing','Security/supportability']], [900,3900,4560])
heading('Code quality findings',2)
bullet('Hard-coded localhost appears across multiple frontend components, defeating the shared Axios abstraction.')
bullet('Duplicate auth helpers, profile components, contribution components and entry points increase drift.')
bullet('Debug println/log levels expose noise and possibly sensitive SQL/data; production logging should be intentional.')
bullet('Generic exception handling and raw entity responses weaken API contracts.')
bullet('No pagination allows unbounded responses; add Pageable and maximum page size.')
bullet('Spring Boot 3.1 and declared libraries should be upgraded deliberately to currently supported, mutually compatible releases; verify migration notes rather than changing versions blindly.')

heading('14. Scaling and HLD evolution',1)
heading('Scale the modular monolith first',2)
num('Make the API stateless and uploads external; run multiple Render instances behind the platform load balancer.')
num('Add connection-pool limits, SQL indexes, pagination and query metrics. Database capacity is often the first bottleneck.')
num('Cache stable semester/subject/resource catalogs with Caffeine or Redis only after measurement; invalidate on admin writes.')
num('Move long-running work such as virus scanning, email and media processing to a queue/worker.')
num('Add CDN/object storage for files; return signed URLs when access control is required.')
heading('When to split services',2)
p('Split only when a module has independent scale, ownership, reliability or release needs. Candidate boundaries might be Identity, Academic Catalog, Community and Assessment. A split adds network failure, distributed tracing, contract/version management, data ownership and deployment overhead.')
code('Phase 1: React -> modular monolith -> PostgreSQL + object storage\nPhase 2: add Redis/cache + worker queue + observability\nPhase 3 (only if justified): API gateway -> identity/catalog/community/assessment services\nEach service owns data; communicate with versioned APIs/events')
heading('Capacity reasoning example',2)
p('Suppose 10,000 monthly students, 1,000 daily active users, peak 50 requests/second and mostly read traffic. A small horizontally replicated Spring service plus indexed PostgreSQL is sufficient. Cache semester catalogs, paginate comments/resources, and keep DB pool sizes bounded. Validate with load tests; do not introduce Kafka or Kubernetes solely for interview appeal.')
heading('Reliability model',2)
bullet('Define SLOs, e.g., 99.9% monthly availability and p95 read latency under 300 ms at target load.')
bullet('Use timeouts for database/external calls, bounded retries only for transient/idempotent operations and circuit breakers only when external dependencies warrant them.')
bullet('Automate backups and perform restore tests; a backup that has never been restored is an assumption.')
bullet('Use rolling/blue-green releases and backwards-compatible migrations: expand schema, deploy code, migrate data, then contract.')

heading('15. Interview preparation',1)
heading('Project ownership questions with model answers',2)
qa = [
('Why did you build this?','Students switch among scattered notes, PYQs, semester information and placement resources. I designed one portal with a normalized academic hierarchy and career/community modules so discovery and progress happen in one product.'),
('Why Spring Boot?','It provided an integrated production-style ecosystem for REST, validation, security and JPA with convention-based setup. The domain is relational and benefits from strong typing. I accepted higher memory/startup cost than a minimal Node service because maintainability and ecosystem integration mattered more.'),
('Why a modular monolith?','The product and team size did not justify distributed systems. One deployable preserved transactions and simple debugging. I separated packages/modules so a future split remains possible if scale or ownership requires it.'),
('Explain login end to end.','React posts credentials. AuthenticationManager delegates to DaoAuthenticationProvider, which loads the database user through UserDetailsService and compares BCrypt hashes. JwtUtil signs a short-lived token. Later Axios sends it; the JWT filter verifies it and populates the SecurityContext before controller authorization.'),
('How is authorization different from authentication?','Authentication proves who the caller is; authorization decides what that identity may do. JWT validation authenticates. Security matchers, @PreAuthorize and resource ownership checks authorize. Client ProtectedRoute is only UX.'),
('Why PostgreSQL rather than MongoDB?','The main entities have strong relationships and integrity rules: branches, semesters, subjects, units, resources, users and owned records. PostgreSQL supports transactions, joins, constraints and useful indexing. MongoDB would be considered if documents were independent and schema flexibility dominated.'),
('What was hardest?','The cross-cutting auth flow and keeping frontend routes, token state, security filters and controller identity consistent. The later audit exposed refresh hydration and authorization-default issues, which taught me to test the full request path, not individual layers only.'),
('How did deployment work?','The intended three-tier deployment put the React build on Vercel, Spring Boot on Render and PostgreSQL on Supabase. The browser calls Render over HTTPS and Render connects to Supabase through JDBC/TLS. Environment variables must carry API URL, CORS origin, JDBC credentials and JWT secret.'),
('What would you improve first?','First security and reproducibility: rotate secrets, default-deny authorization, enforce ownership, environment-driven config, Flyway migrations and auth integration tests. Then durable object storage, unified frontend API/auth handling, pagination and observability.'),
('What design patterns did you use?','Layered architecture, Repository, dependency injection, DTO, Spring Security filter chain/chain of responsibility, strategies such as PasswordEncoder and AuthenticationProvider, React observer-style rendering, and a polymorphic context association for comments.'),
('How would you scale it?','Keep the API stateless, externalize files, add indexes/pagination and measure queries, then horizontally replicate. Cache read-heavy catalogs and add workers for asynchronous jobs. Split services only when independent scaling or ownership justifies distributed-system cost.'),
('How do you avoid N+1 queries?','Observe SQL, keep relationships lazy by default, fetch explicit API shapes using projections/entity graphs or fetch joins, batch where appropriate, and verify query counts in integration tests.'),
]
for q,a in qa:
    heading(q,3); p(a)
heading('Rapid-fire technical questions',2)
table([['Question','Answer cue'],['What does @SpringBootApplication combine?','@Configuration, @EnableAutoConfiguration and @ComponentScan.'],['Why OncePerRequestFilter?','Guarantees one filter execution per request dispatch pattern; suitable for token extraction.'],['401 vs 403?','401: unauthenticated/invalid credentials; 403: authenticated but not permitted.'],['Why BCrypt?','Salted, adaptive one-way hash; work factor slows brute force.'],['What is a Spring bean scope concern?','Singleton services serve concurrent requests; keep mutable request state out of fields.'],['What is JPA dirty checking?','Managed entity changes are detected and flushed inside a transaction.'],['PUT vs PATCH?','PUT replaces a resource representation; PATCH applies a partial change.'],['Why DTO validation?','Reject malformed input at the boundary and avoid persistence/domain leakage.'],['Why pagination?','Bounds latency/memory/payload and stabilizes APIs as data grows.'],['Why exact CORS origins?','CORS is browser policy; least-privilege origins reduce unwanted cross-site calls.'],['Why migrations?','Versioned, repeatable, reviewable schema evolution across environments.'],['Can JWT be invalidated?','Not inherently before expiry; use short TTL/refresh rotation or a deny-list when needed.']], [3600,5760])
heading('Honest answers for gaps',2)
callout('Use this wording.', '“The repository compiles, but it currently has no backend tests and deployment config has drifted back to localhost. I would not claim those are production-ready. I documented the gap, reproduced the build, and prioritized security/configuration/integration tests before new features.”', NAVY, 'EEF3F8')

heading('16. Learning plan and cheat sheets',1)
heading('Seven-day mastery plan',2)
table([['Day','Hands-on task','Speak aloud'],['1','Run locally; draw feature/repository map','90-second product and stack story'],['2','Debug register/login/filter with breakpoints','Authentication vs authorization; JWT trade-offs'],['3','Trace semester -> subject -> unit -> resource SQL','JPA relationships, lazy loading, N+1'],['4','Trace career, quiz, contribution and comments','Business rules and DTO boundaries'],['5','Trace routes, Zustand, Axios and refresh behavior','React state and API contract'],['6','Reconstruct Vercel/Render/Supabase config in a safe branch','Deployment, CORS, secrets, migrations'],['7','Write top tests and do mock interview','Improvements, scale, failures, ownership']], [800,4100,4460])
heading('Breakpoint map',2)
bullet('AuthController.login/register -> CustomUserDetailsService -> JwtUtil.generateToken.')
bullet('JwtAuthenticationFilter.doFilterInternal for any authenticated call.')
bullet('StudyResourceController subject endpoint -> StudyResourceService -> SubjectRepository.')
bullet('ContributionController add and CommentController create to inspect Authentication and entity ownership.')
bullet('QuizController submit to observe server-side scoring.')
bullet('Axios request interceptor and Zustand setUser/logout to inspect browser auth state.')
heading('Before the interview checklist',2)
for x in ['Remove/rotate real secrets and know why','Know five endpoint flows without notes','Draw the three-tier HLD in under two minutes','Explain four technology trade-offs, not slogans','Name three implemented patterns and two patterns not present','Explain the biggest security bug and exact fix','Know build/test status honestly','Prepare one failure story and one refactoring story','Open deployed URLs and verify health if they still exist','Quantify usage only if you have real measurements']:
    bullet(x)
heading('Final one-page mental model',2)
code('PRODUCT: academic discovery + career prep + community/productivity\nCLIENT: React routes/components -> Zustand auth -> Axios\nSECURITY: BCrypt login -> HS256 JWT -> filter -> SecurityContext -> authorization\nSERVER: Controller -> Service -> Repository -> JPA/Hibernate\nDATA: PostgreSQL relational hierarchy and user-owned records\nDEPLOY: Vercel SPA -> Render API -> Supabase Postgres\nCURRENT RISKS: secrets, permitAll, localhost/CORS, ephemeral uploads, no migrations/tests\nNEXT: secure config -> default deny/ownership -> tests -> Flyway -> storage -> observability')

heading('17. Future product and AI integration blueprint',1)
callout('North-star product.', 'Evolve the portal from a resource directory into a trusted, personalized learning operating system: it should understand the curriculum, student goals and progress; recommend the next best action; answer from approved material with citations; and help educators improve content without replacing academic judgment.', NAVY, 'EEF3F8')
heading('17.1 Product evolution principles',2)
bullet('Fix the platform foundation before adding AI: security, environment configuration, migrations, tests, durable storage, observability and stable API contracts are prerequisites.')
bullet('Start with high-frequency student problems, not with a chatbot. Every AI feature should reduce search time, improve comprehension or create a measurable study action.')
bullet('Ground academic answers in institution-approved resources and display citations. A fluent answer without evidence is not trustworthy enough for exam preparation.')
bullet('Keep humans in control for grades, moderation, career decisions and published content. AI may recommend or draft; authorized users approve consequential actions.')
bullet('Design graceful fallback. If the model, vector store or network is unavailable, normal browsing, todos, quizzes and progress tracking must continue to work.')
heading('17.2 Non-AI product roadmap',2)
table([['Area','Better solution','Business/student impact'],['Academic catalog','Admin CMS, versioned syllabus, resource moderation, search, tags and quality scores','Reliable and discoverable material'],['Learning progress','Per-resource completion, time tracking, streak rules and mastery by topic','Visible momentum and personalized next steps'],['Assessment','Question bank, timed tests, explanations, attempt history and analytics','Exam readiness based on evidence'],['Collaboration','Moderated discussions, accepted answers, reports, notifications and reputation','Safer peer learning'],['Planning','Calendar sync, deadlines, timetable, reminders and weekly review','One actionable student workspace'],['Career','Role roadmaps, skill-gap analysis, resume/projects and interview practice','Connect coursework with outcomes'],['Accessibility','Keyboard navigation, captions, readable contrast, screen-reader semantics and multilingual UI','Inclusive learning'],['Administration','Role-based dashboards, audit logs, content workflow and usage analytics','Operational control and accountability']], [1500,4100,3760])
heading('17.3 Recommended target architecture',2)
code('Browser / mobile-ready PWA\n  -> Vercel CDN: React application\n  -> API boundary: Render Spring Boot modular monolith\n       |-- Identity and RBAC\n       |-- Academic Catalog and Search\n       |-- Progress and Assessment\n       |-- Community and Moderation\n       |-- AI Orchestration module\n       |      |-- prompt templates and policy checks\n       |      |-- retrieval and citations\n       |      |-- model-provider adapter\n       |      `-- evaluation/feedback logging\n       |-- background jobs / queue\n       v\n  Supabase PostgreSQL + pgvector | Supabase Storage | Redis/cache (when justified)\n       |\n  Approved LLM and embedding providers through server-side calls only')
p('Keep this as one deployable initially. Create internal module interfaces and separate database schemas/tables by domain. Extract an AI worker or service only when asynchronous load, GPU/provider requirements, independent releases or fault isolation justify it.')
heading('17.4 Highest-value AI capabilities',2)
table([['Capability','User flow','AI technique','First-version guardrail'],['Cited study assistant','Ask about a subject/unit and receive an answer linked to notes/PYQs','Retrieval-augmented generation (RAG)','Only retrieved approved sources; answer “not found” when evidence is weak'],['Semantic search','Search by meaning instead of exact keywords','Embeddings + vector similarity + metadata filters','Always filter by branch, semester, subject and access rights'],['Personalized study plan','Goal/date/availability -> weekly tasks','Rules first, LLM for explanation and plan wording','User confirms plan; deterministic deadline and workload constraints'],['Quiz generation','Generate practice from selected material','RAG + structured JSON generation','Teacher/admin approval for published questions; source citation and duplicate checks'],['Answer explanation','Explain why an option is correct and diagnose misconception','Grounded generation using question/source context','Never change official answer silently; show confidence/evidence'],['Notes summarization','Produce revision notes, flashcards or concept map','Chunk-aware summarization','Label as generated; retain source links; user controls length/level'],['Career coach','Compare current skills with a target role and propose projects','Profile + curated role taxonomy + retrieval','No guarantee of employment; avoid sensitive-trait inference'],['Interview practice','Role-specific mock interview with rubric and feedback','Conversation model + rubric-based evaluator','Store only with consent; show rubric and allow deletion'],['Content moderation','Prioritize spam, abuse or unsafe uploads for review','Classifier + rules','Human decides removal/discipline; record reasons and appeals'],['Admin insight assistant','Natural-language analytics over approved aggregates','Text-to-query through constrained semantic layer','Read-only queries, allow-listed metrics, row limits and audit logs']], [1450,2500,2700,2710])
heading('17.5 RAG: the recommended first AI project',2)
p('RAG is a better first step than fine-tuning because the portal’s knowledge changes with notes, syllabi and contributions. Retrieval supplies current project-owned context at request time, supports citations and can enforce course/access filters. Fine-tuning is more suitable later for stable behavior, tone or classification patterns; it is not a reliable way to inject frequently changing facts.')
heading('Ingestion pipeline',3)
num('An authorized uploader stores the original PDF/document in Supabase Storage and creates a resource record in PostgreSQL.')
num('A background job validates file type, scans for malware, extracts text/OCR, normalizes it and preserves page/section metadata.')
num('The pipeline splits text into overlapping semantic chunks. Each chunk stores resource_id, subject_id, unit_id, page, heading, version, visibility and checksum.')
num('An embedding provider converts each chunk into a vector. Store it in PostgreSQL with pgvector or a dedicated vector database only if scale demands it.')
num('Mark the resource searchable only after extraction and indexing succeed. Re-index changed versions and retire old chunks safely.')
heading('Question-answer pipeline',3)
code('Question + authenticated user\n  -> authorize subject/resource access\n  -> normalize query and derive metadata filters\n  -> embed query\n  -> vector + keyword hybrid retrieval\n  -> optional reranker\n  -> top evidence chunks with page metadata\n  -> prompt model to answer only from evidence\n  -> validate structured response and citations\n  -> return answer + sources + “insufficient evidence” state\n  -> collect explicit helpful/not-helpful feedback')
heading('Suggested AI data model',3)
table([['Table','Important fields','Purpose'],['document_version','id, resource_id, storage_path, checksum, status, created_at','Immutable source/version tracking'],['document_chunk','id, version_id, subject_id, unit_id, page, text, embedding, token_count','Retrievable evidence with metadata'],['ai_conversation','id, user_id, subject_id, title, consent/retention fields','Conversation ownership and deletion'],['ai_message','id, conversation_id, role, content, model, latency, token counts','Auditable interaction metadata'],['ai_citation','message_id, chunk_id, rank, score','Evidence trace for each answer'],['ai_feedback','message_id, user_id, rating, reason','Quality signal and evaluation dataset'],['ai_job','id, type, status, attempts, error, timestamps','Reliable asynchronous ingestion/generation'],['prompt_version','feature, version, template hash, active flag','Reproducible prompt releases']], [1700,3300,4360])
heading('17.6 AI service design in Spring Boot',2)
p('Add an ai package/module rather than calling a model directly from controllers. This isolates provider APIs, prompt versions, retrieval, safety and observability. Use constructor injection and interfaces so the provider can change without rewriting business logic.')
code('ai/\n  controller/AiTutorController\n  service/AiTutorService, RetrievalService, IngestionService\n  port/LanguageModelClient, EmbeddingClient, VectorSearchPort\n  adapter/<provider-specific clients>\n  policy/AiAccessPolicy, ContentSafetyPolicy\n  dto/AiQuestionRequest, GroundedAnswerResponse\n  evaluation/AiEvaluationService')
code('public interface LanguageModelClient {\n  GroundedCompletion complete(GroundedPrompt prompt);\n}\n\npublic record GroundedAnswerResponse(\n  String answer, List<CitationDTO> citations,\n  boolean insufficientEvidence, String requestId) {}')
bullet('Controller: authenticate, validate and rate-limit; never accept an arbitrary system prompt or model name from the browser.')
bullet('Orchestration service: retrieve evidence, construct a versioned prompt, call the provider, validate output and persist permitted telemetry.')
bullet('Provider adapter: timeouts, bounded retries, error mapping, token/cost accounting and circuit breaking where appropriate.')
bullet('Background worker: ingestion, OCR, embeddings, large quiz generation and re-indexing; do not hold an HTTP request open for minutes.')
heading('17.7 Personalization without becoming invasive',2)
p('Start with explicit and educational signals: branch, semester, selected goal, deadlines, completed resources, quiz mastery and preferred study time. Avoid inferring health, caste, religion, socioeconomic status or other sensitive characteristics. Let users inspect, correct, disable and delete personalization data.')
code('candidate resources\n  -> authorization + curriculum filters\n  -> mastery gap score\n  -> deadline urgency\n  -> prerequisite readiness\n  -> quality/freshness\n  -> diversity constraint\n  -> ranked recommendations with a human-readable “Why this?”')
bullet('Begin with a deterministic weighted recommender. It is explainable, testable and works before sufficient behavioral data exists.')
bullet('Later add collaborative/content-based ranking only after consent, data-quality checks, bias analysis and offline/online evaluation.')
bullet('Never optimize solely for time-on-platform. Optimize for learning outcomes, plan completion and user-reported usefulness.')
heading('17.8 AI safety, privacy and security',2)
table([['Risk','Example','Control'],['Hallucination','Tutor invents a formula or deadline','RAG, citations, evidence threshold, abstention and feedback'],['Prompt injection','Uploaded note says to ignore policy or leak data','Treat retrieved text as untrusted data; fixed system policy; tool allow-list; output validation'],['Data leakage','One user retrieves private contribution/chat','Authorization before retrieval and metadata filters in every query'],['PII exposure','Resume/profile sent unnecessarily to a provider','Data minimization, redaction, consent, provider retention settings and deletion workflow'],['Unsafe advice','Overconfident mental-health, financial or career claims','Scope boundaries, safety classifiers, crisis/help routing and clear limitations'],['Bias','Career recommendations disadvantage groups','Avoid sensitive proxies, evaluate cohorts, explain recommendations and allow correction'],['Cost abuse','Automated requests exhaust model budget','Per-user quotas, rate limits, caching, max tokens and budget alerts'],['Model/vendor outage','Tutor becomes unavailable','Timeouts, graceful fallback to search, provider abstraction and status message'],['Copyright/provenance','Generated notes reproduce protected material','Use licensed/authorized content, citations, access rules and output-length controls'],['Evaluation gaming','Model scores well on easy synthetic cases','Human-authored golden set, adversarial tests and production feedback sampling']], [1700,2700,4960])
heading('17.9 Evaluation and observability',2)
p('AI quality is not proven by a successful API response. Build an evaluation set from real, permissioned subject questions with expected evidence and rubrics. Run it whenever prompts, retrieval, embeddings or models change.')
table([['Layer','Metrics'],['Retrieval','Recall@k, precision@k, mean reciprocal rank, correct-source rate, permission-filter violations'],['Answer quality','Groundedness, citation correctness, completeness, factual error rate, abstention quality'],['Product','Helpful rating, task completion, repeated-question rate, study-plan adherence, quiz improvement'],['Operations','p50/p95 latency, availability, timeout/error rate, queue age, tokens and cost per feature/user'],['Safety/privacy','Prompt-injection success rate, unsafe-output rate, PII incidents, cross-user leakage tests']], [1800,7560])
bullet('Log request IDs, prompt version, retrieval chunk IDs, model version, latency and token counts; avoid logging raw private content by default.')
bullet('Use online A/B experiments only after offline safety gates. Define a rollback threshold before launch.')
bullet('Maintain a small human review sample and a user feedback path. Automated “LLM-as-judge” scores are useful signals, not ground truth.')
heading('17.10 Phased implementation roadmap',2)
table([['Phase','Deliverable','Exit criteria'],['0. Foundation','Security fixes, unified config/auth, Flyway, tests, storage, metrics','Stable staging; ownership tests; no tracked secrets'],['1. Better portal','Search/tags, progress, moderation, notifications and analytics','Core journeys reliable; measurable baseline'],['2. AI pilot','Cited Q&A for one subject using approved PDFs','High citation correctness; zero permission leaks; useful fallback'],['3. Learning tools','Summaries, flashcards and quiz drafts from selected sources','Teacher/user review workflow; quality thresholds'],['4. Personalization','Rules-based study planner and explainable recommendations','Improved plan completion without harmful cohort disparity'],['5. Scale and platform','Queue/workers, caching, model routing, cost controls and broader subjects','SLOs and budget met under load'],['6. Advanced experiences','Multilingual tutor, voice/accessibility, interview simulations','Consent, safety and evaluation gates passed']], [1200,4800,3360])
heading('A realistic first 30-day build',3)
num('Week 1: secure the existing application, introduce environment profiles/Flyway and write authentication/ownership integration tests.')
num('Week 2: move uploads to Supabase Storage; add document_version and ai_job; implement safe PDF text extraction in a worker.')
num('Week 3: enable pgvector, chunk one approved subject, implement metadata-filtered retrieval and citations.')
num('Week 4: add a subject-scoped tutor UI, rate limiting, feedback, evaluation dataset and staging rollout to a small consenting group.')
heading('17.11 Cost and performance controls',2)
bullet('Estimate cost before launch: questions per active user x retrieved tokens + generated tokens x provider price, plus embedding, storage and database costs.')
bullet('Cache embeddings permanently by checksum and cache only safe, non-personal answers with source-version-aware keys.')
bullet('Use a small model for classification/query rewriting and a stronger model only when complexity or evidence warrants it.')
bullet('Cap context, output tokens, conversation history and requests per minute/day. Summarize old conversation turns if retention is allowed.')
bullet('Batch embeddings and asynchronous work; monitor cost per successful learning task rather than cost per raw request alone.')
heading('17.12 Future integrations',2)
table([['Integration','Use','Design note'],['Learning Management System','Import courses, assignments and grades','Prefer standards/APIs such as LTI/OneRoster where supported; explicit permissions'],['Google/Microsoft Calendar','Deadlines, study blocks and reminders','OAuth scopes should be minimal; support disconnect and token revocation'],['Email/push notifications','Due dates, plan changes and moderation','User preferences, quiet hours and unsubscribe controls'],['Supabase Storage/CDN','Durable notes, resumes and media','Private buckets, signed URLs, scanning and retention'],['Video platforms','Curated lectures and transcripts','Store metadata/links; respect API terms and copyright'],['College SSO','Single institutional login','OIDC/SAML via a mature identity solution; map roles carefully'],['Coding judge','Run programming assessments','Isolated sandbox, resource limits and untrusted-code security'],['Analytics warehouse','Long-term product/learning analytics','Pseudonymize, minimize and separate operational from analytical access']], [1900,3300,4160])
heading('17.13 AI interview questions and model answers',2)
aiqa = [
('Where would you add AI first?','I would start with subject-scoped semantic search and cited Q&A over approved resources. It solves a real discovery problem, is measurable, and RAG gives provenance. I would not begin with an unrestricted general chatbot.'),
('Why RAG rather than fine-tuning?','The knowledge changes and must be cited. RAG retrieves current authorized material per request. Fine-tuning is useful later for stable behavior or format, but not as the source of frequently changing course facts.'),
('How do you prevent hallucinations?','Retrieve authorized evidence, require the model to answer only from it, return citations, set an evidence threshold, support abstention, validate structured output and continuously evaluate groundedness and citation correctness.'),
('How do you prevent one user seeing another user’s data?','Authorize before retrieval and carry tenant/user/visibility filters into both SQL and vector search. Never retrieve broadly and filter after generation. Add adversarial integration tests for cross-user leakage.'),
('Would AI be a microservice?','Not initially. I would use an internal AI module plus background worker in the modular monolith. I would extract it only for independent scaling, release cadence, provider/GPU isolation or fault containment.'),
('How would you evaluate the tutor?','Use a human-authored golden set with expected source passages and rubrics, measuring retrieval recall, groundedness, citation correctness, factual error and abstention. Then track latency, cost, helpfulness and learning outcomes in a controlled rollout.'),
('How do you control model cost?','Quotas, token limits, model routing, batching/caching embeddings, safe response caching, concise prompts, asynchronous jobs and cost dashboards. The useful metric is cost per successful learning task.'),
('What happens if the AI provider fails?','Timeout quickly, avoid unsafe retry storms, show a transparent unavailable state and fall back to normal keyword/resource search. Core portal functions must not depend on the model.'),
('How would personalization work with little data?','Start with explicit profile and deterministic rules: curriculum, deadlines, prerequisites, mastery and availability. This avoids cold-start problems and remains explainable. Add learned ranking only after enough consented high-quality data exists.'),
('What is prompt injection here?','Malicious text in a user question or uploaded document attempts to override instructions or access tools/data. Retrieved content is treated as untrusted, system policy is fixed, tools are allow-listed, access is enforced outside the model and outputs are validated.'),
]
for q,a in aiqa:
    heading(q,3); p(a)
heading('17.14 What not to build yet',2)
bullet('Do not introduce microservices, Kafka, Kubernetes or a separate vector database before measured scale or ownership requires them.')
bullet('Do not let a model directly execute arbitrary SQL, browse private files or call unrestricted tools.')
bullet('Do not auto-grade high-stakes exams, delete community content or make career decisions without human review and appeal.')
bullet('Do not train on student conversations, resumes or contributions without explicit lawful consent and retention controls.')
bullet('Do not promise “100% accurate AI.” Define measurable thresholds, show evidence and design for uncertainty.')
heading('17.15 Future-state summary',2)
code('FOUNDATION: secure, tested modular monolith + migrations + durable storage\nDATA: versioned approved content + progress + explicit consent\nAI: retrieval first -> grounded generation -> citations -> feedback/evaluation\nPERSONALIZATION: explainable rules first -> learned ranking only with evidence\nOPERATIONS: async jobs, quotas, model routing, observability and fallback\nGOVERNANCE: least privilege, human oversight, deletion, audit and safety testing\nOUTCOME: a trustworthy learning assistant, not merely an AI chat box')

heading('Appendix A. Endpoint reference',1)
table([['Area','Endpoints present in source'],['Auth','POST /api/auth/login; POST /api/auth/register; GET /api/auth/test'],['Users/admin','GET /api/users/all; GET /api/users/profile; GET /api/admin/users; PUT /api/internal/promote-to-admin'],['Academics','GET /api/semesters; GET /api/semesters/{semester}/subjects; GET /api/subjects/{id}/units; GET /api/units/{id}; GET /api/units/{id}/resources; GET /api/units/{id}/pyqs'],['Academic admin','POST /api/admin/semesters; /subjects; /units; /resources/upload'],['Career','GET /api/career/{goal}; GET /api/career/{goal}/{type}; POST /api/career/admin/upload; GET /api/career/placement/progress; POST /api/career/placement/attempt; GET /api/career/goal'],['Quiz','GET /api/quiz/{goal}; POST /api/quiz/submit'],['Contributions','POST /api/contributions/add; GET /api/contributions/my; GET /{id}/upvotes, /downvotes, /bookmarks'],['Comments','POST/GET /api/comments/{type}/{id}; GET /api/comments/user'],['Todos','GET /api/todos/test; GET /api/todos/{userId}; POST /api/todos; PUT/DELETE /api/todos/{id}'],['Diagnostics','GET /api/test/hello; GET /api/test/admin; GET /']], [1600,7760])
callout('Caution.', 'An endpoint being present does not mean it is correctly protected, consumed by the frontend or production-ready. Use Chapters 6, 8 and 13 alongside this list.', RED, 'FCEEEE')
heading('Appendix B. Source files to read in order',1)
for x in ['backend/build.gradle and application.properties','StudentCompanionApplication.java','config/SecurityConfig.java','filter/JwtAuthenticationFilter.java and util/JwtUtil.java','controller/AuthController.java and service/CustomUserDetailsService.java','model/User.java, Role.java and UserRepository.java','StudyResourceController.java -> StudyResourceService.java -> academic repositories/entities','CareerController.java, QuizController.java, ContributionController.java, CommentController.java, TodoController.java','frontend/package.json, src/index.js, src/App.js','src/stores/authStore.js and src/utils/axiosConfig.js','LoginPage, SignupPage, ProtectedRoute','SubjectsPage -> UnitPage; CareerGoalPage; ContributePage','database/schema.sql only after understanding that it exceeds implemented entities']:
    num(x)
heading('Appendix C. Review provenance',1)
p('This guide was derived from the checked-in Java, JavaScript/JSX, Gradle, npm, SQL and properties files in the repository. The backend was compiled through its Gradle wrapper. No external production dashboard was accessed, and no historical deployment environment values were available. Recommendations are labeled as targets rather than existing implementation.')

# prevent row splitting and repeat headers
for t in doc.tables:
    trPr=t.rows[0]._tr.get_or_add_trPr(); hdr=OxmlElement('w:tblHeader'); hdr.set(qn('w:val'),'true'); trPr.append(hdr)
    for row in t.rows:
        pr=row._tr.get_or_add_trPr(); cant=OxmlElement('w:cantSplit'); pr.append(cant)

doc.core_properties.title='Student Study Portal - Complete Learning and Interview Guide'
doc.core_properties.subject='Codebase walkthrough, HLD, system design, deployment and interview preparation'
doc.core_properties.author='Codex - repository-derived technical documentation'
doc.save(OUT)
print(OUT)
